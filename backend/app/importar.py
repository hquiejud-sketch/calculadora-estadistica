"""Lectura de datos desde archivos CSV, TXT o Excel."""

import io
import re

NUMERO = re.compile(r"-?\d+(?:\.\d+)?")


def numeros_de_texto(texto: str) -> list[float]:
    # Si usan coma decimal (1,5;2,7) y punto y coma como separador, se normaliza.
    if ";" in texto and re.search(r"\d,\d", texto):
        texto = texto.replace(",", ".")
    return [float(n) if "." in n else int(n) for n in NUMERO.findall(texto)]


def leer_archivo(nombre: str, contenido: bytes) -> list[float]:
    nombre = nombre.lower()
    if nombre.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        datos = []
        for fila in libro.active.iter_rows(values_only=True):
            for celda in fila:
                if isinstance(celda, bool):
                    continue
                if isinstance(celda, (int, float)):
                    datos.append(celda)
                elif isinstance(celda, str):
                    datos.extend(numeros_de_texto(celda))
        return datos
    if nombre.endswith((".csv", ".txt", ".tsv")):
        return numeros_de_texto(contenido.decode("utf-8", errors="ignore"))
    raise ValueError("Formato no soportado. Usa .csv, .txt o .xlsx")
